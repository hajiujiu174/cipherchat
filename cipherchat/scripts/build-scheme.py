"""Create the formal report from the retained Word template; render separately with Word COM."""
from pathlib import Path
import hashlib
import json
import zipfile
from datetime import datetime
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT / 'build/document-work'
TEMPLATE = WORK / 'template-reference.docx'
OUT = ROOT / 'artifacts/CipherChat_项目方案书.docx'
BLACK = RGBColor(0, 0, 0)
WIDTH = 14.663
WORK.mkdir(parents=True, exist_ok=True)


def diagram():
    image = Image.new('RGB', (1680, 790), 'white')
    d = ImageDraw.Draw(image)
    font = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 30)
    small = ImageFont.truetype('C:/Windows/Fonts/msyh.ttc', 25)
    def box(rect, title, lines):
        d.rounded_rectangle(rect, radius=14, fill='#F4F5F6', outline='#30363B', width=3)
        x1,y1,x2,y2 = rect
        d.text(((x1+x2)/2,y1+27), title, font=font, fill='black', anchor='mt')
        for index, line in enumerate(lines):
            d.text(((x1+x2)/2,y1+85+index*42), line, font=small, fill='#222222', anchor='mt')
    def arrow(start, end, label=None):
        d.line((start,end), fill='#30363B', width=4)
        x,y=end
        if end[0] == start[0]: d.polygon([(x,y),(x-11,y-18),(x+11,y-18)], fill='#30363B')
        else: d.polygon([(x,y),(x-18,y-11),(x-18,y+11)], fill='#30363B')
        if label: d.text(((start[0]+end[0])/2,start[1]-34), label, font=small, fill='black', anchor='mt')
    box((30,25,750,180),'Qt 主线程',['账号　私聊　文件　群聊  |  命令 / 事件队列'])
    box((30,245,750,455),'单一业务线程',['ClientApplication 编排服务','账号 / 会话 / 聊天 / 文件 / 群聊','SQLite 原子快照 + Windows DPAPI'])
    box((30,525,750,740),'客户端网络线程',['TLS 1.3 验证与帧收发','SSL 对象由同一网络线程拥有','有界队列向业务线程交付消息'])
    box((1020,245,1640,455),'服务端',['账号认证 / 预密钥 / 群成员校验','持久化密文队列 / 路由 / 签收','持有 TLS 私钥与路由元数据'])
    box((1020,525,1640,740),'另一客户端',['独立身份与会话状态','本地解密 / 校验 / 提交后签收','服务器不持有双方会话私钥'])
    arrow((390,180),(390,245)); arrow((390,455),(390,525))
    arrow((1320,455),(1320,525))
    d.line([(750,585),(885,585),(885,350),(1020,350)],fill='#30363B',width=4)
    d.polygon([(1020,350),(1002,339),(1002,361)],fill='#30363B')
    d.text((885,260),'TLS 中转',font=small,fill='black',anchor='mt')
    d.text((885,301),'内含端到端密文',font=small,fill='black',anchor='mt')
    path = WORK / 'architecture.png'; image.save(path); return path


doc = Document(TEMPLATE)
body = doc._element.body
for child in list(body):
    if child.tag != qn('w:sectPr'): body.remove(child)
for style in doc.styles:
    if style.type == WD_STYLE_TYPE.PARAGRAPH:
        props = style.element.find(qn('w:pPr'))
        if props is not None:
            numbering = props.find(qn('w:numPr'))
            if numbering is not None: props.remove(numbering)


def fonts(style, size, bold=False, western='Times New Roman'):
    style.font.name = western; style.font.size = Pt(size)
    style.font.bold = bold; style.font.color.rgb = BLACK
    rf = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
        rf.attrib.pop(qn('w:'+attribute), None)
    for attribute,value in [('ascii',western),('hAnsi',western),('eastAsia','宋体'),('cs',western)]: rf.set(qn('w:'+attribute),value)


def style(name, size, bold=False, indent=0, before=0, after=0, spacing=1.25):
    s = doc.styles[name] if name in doc.styles else doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    fonts(s,size,bold)
    p = s.paragraph_format; p.first_line_indent = Pt(indent)
    p.left_indent = Pt(0); p.right_indent = Pt(0)
    p.space_before = Pt(before); p.space_after = Pt(after); p.line_spacing = spacing
    p.keep_together = False; p.widow_control = True
    snap=OxmlElement('w:snapToGrid'); snap.set(qn('w:val'),'0'); s.element.get_or_add_pPr().append(snap)
    return s


style('CC Body',12,indent=24,after=4)
style('Title',24,True,after=20)
for level,size,before,after in [(1,16,8,8),(2,14,7,3),(3,12,3,3)]:
    s = style(f'CC Heading {level}',size,True,before=before,after=after)
    s.paragraph_format.keep_with_next=True
    outline = OxmlElement('w:outlineLvl'); outline.set(qn('w:val'),str(level-1)); s.element.get_or_add_pPr().append(outline)
style('CC Caption',10.5,after=7,spacing=1.05)
style('CC Table',10.5,spacing=1.10)
style('CC Reference',10.5,after=5,spacing=1.15)
style('CC Code',9,after=3,spacing=1.12)
fonts(doc.styles['CC Code'],9,western='Consolas')
style('CC Spacer',1,after=5,spacing=1)
style('CC Cover',16,after=18,spacing=1.4)
for section in doc.sections:
    section.page_width=Cm(21.0026); section.page_height=Cm(29.7039)
    section.top_margin=Cm(3.1397); section.bottom_margin=Cm(2.54)
    section.left_margin=Cm(3.1697); section.right_margin=Cm(3.1697)
    section.header_distance=Cm(2.05); section.footer_distance=Cm(1.75)
    section.different_first_page_header_footer=True
    for part in (section.header,section.first_page_header,section.footer,section.first_page_footer):
        for p in list(part.paragraphs)[1:]: p._element.getparent().remove(p._element)
        part.paragraphs[0].clear()
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run('- ')
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
    footer.add_run(' -')


def p(text, style_name='CC Body'):
    par=doc.add_paragraph(text,style_name)
    if style_name=='CC Body': par.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    return par


def h(text, level=1, page=False):
    par=p(text,f'CC Heading {level}')
    par.alignment=WD_ALIGN_PARAGRAPH.CENTER if level==1 else WD_ALIGN_PARAGRAPH.LEFT
    par.paragraph_format.page_break_before=page
    return par


def figure(path, caption, width=WIDTH):
    par=p('','CC Caption'); par.paragraph_format.keep_with_next=True; par.alignment=WD_ALIGN_PARAGRAPH.CENTER
    shape=par.add_run().add_picture(str(path),width=Cm(width))
    shape._inline.docPr.set('descr',caption)
    cap=p(caption,'CC Caption'); cap.alignment=WD_ALIGN_PARAGRAPH.CENTER


def table(caption, headers, rows, widths, short=()):
    par=p(caption,'CC Caption'); par.alignment=WD_ALIGN_PARAGRAPH.CENTER; par.paragraph_format.keep_with_next=True
    t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    pr=t._tbl.tblPr
    tw=pr.find(qn('w:tblW')); tw.set(qn('w:type'),'dxa'); tw.set(qn('w:w'),str(round(sum(widths)*567)))
    margins=OxmlElement('w:tblCellMar')
    for tag,amount in [('top',85),('left',95),('bottom',85),('right',95)]:
        x=OxmlElement('w:'+tag); x.set(qn('w:w'),str(amount)); x.set(qn('w:type'),'dxa'); margins.append(x)
    pr.append(margins)
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        x=OxmlElement('w:'+edge); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'777777'); borders.append(x)
    pr.append(borders)
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    for index, row in enumerate([headers]+rows):
        cells=t.rows[0].cells if index==0 else t.add_row().cells
        trpr=cells[0]._tc.getparent().get_or_add_trPr(); trpr.append(OxmlElement('w:cantSplit'))
        if index==0: trpr.append(OxmlElement('w:tblHeader'))
        for i,(cell,text) in enumerate(zip(cells,row)):
            cell.width=Cm(widths[i]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            par=cell.paragraphs[0]; par.style=doc.styles['CC Table']; par.add_run(str(text))
            par.alignment=WD_ALIGN_PARAGRAPH.CENTER if index==0 or i in short else WD_ALIGN_PARAGRAPH.LEFT
            if index==0:
                par.paragraph_format.keep_with_next=True
                for run in par.runs: run.bold=True
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'EEEEEE'); cell._tc.get_or_add_tcPr().append(shade)
    p('','CC Spacer')
    return t


# 1 Cover
cover=p('课程设计项目方案书','CC Cover'); cover.alignment=WD_ALIGN_PARAGRAPH.CENTER
cover.paragraph_format.space_before=Pt(85)
title=p('CipherChat','Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
title=p('端到端加密聊天系统','Title'); title.alignment=WD_ALIGN_PARAGRAPH.CENTER
subtitle=p('设计方案与实现验收','CC Cover'); subtitle.alignment=WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after=Pt(95)
for line in ['Windows 桌面客户端与密文中转服务端','C++20  ·  Qt 6  ·  OpenSSL 3  ·  SQLite','2026 年 9 月 14 日']:
    par=p(line,'CC Cover'); par.alignment=WD_ALIGN_PARAGRAPH.CENTER

# 2 Abstract and requirements
h('摘 要',page=True)
p('CipherChat 面向课程设计中的双端安全通信与可恢复传输场景，实现了 Windows 图形客户端、TLS 服务端和自动化验收工程。系统以账号认证和显式身份核对建立信任，使用 X3DH 初始化会话，通过 Double Ratchet 加密私聊及文件元数据；聊天签收、群发结果与文件检查点均纳入持久化流程。')
p('本次实现将原有密码算法接入统一业务服务，修复了 TLS 并发读写、认证失败推进棘轮、离线消息过早删除及短文件块拒收等问题。纯核心验收、双原生窗口演示、跨进程恢复和发布包启动均有独立证据。该协议是自定义的课程设计简化协议，不与 Signal 互通，功能测试不构成完整安全证明。')
par=p('关键词：端到端加密；身份核对；原子快照；文件续传；Qt 桌面应用'); par.paragraph_format.first_line_indent=Pt(0)
h('1 需求与验收目标')
table('表1.1 功能范围与主要验收点',['能力','实现要求','验收任务'],[
    ['账号与信任','验证 TLS 身份；注册、登录、预密钥与人工指纹核对','M1-03，M2-01/02，M3-02'],
    ['可靠私聊','UTF-8 文本加密、路由绑定、去重、双方签收','M2-03，M4-01'],
    ['文件传输','64 KiB 分块；短文件、乱序、篡改和跨进程恢复','M4-02/03'],
    ['群聊与状态','逐成员加密、成员授权、在线与输入提示','M5-01/02'],
    ['桌面与交付','四页 GUI、单业务线程、恢复、发布包和方案书','M3-05，M4-04，M5-03，M6-01'],
], [2.4,8.0,4.263], short=(0,))
p('验收按 M0—M6 共 22 项任务登记。完成条件以同名业务 CTest、实际进程运行和界面证据为准；仅能编译或通过接口契约测试的模块不计为业务完成。')

# 3 Architecture
h('2 系统架构',page=True)
p('系统分为公共协议与密码算法、客户端业务服务、Qt 界面和服务端四部分。公共算法只保留一份实现；正式客户端以 ClientApplication 为组合根，不创建旧 ClientCore 处理同一连接。图2-1说明对象与线程的职责。')
figure(diagram(),'图2-1 线程职责与加密通信路径')
h('2.1 客户端分层',2)
p('Qt 主线程负责输入与显示，命令进入有界队列；业务线程串行调用账号、会话、聊天、文件和群服务，同时拥有 SQLite 与会话状态。每次 tick 最多处理 64 条网络消息，文件操作也按块推进，避免一次读完整文件。网络线程仅做连接、TLS 收发、帧解析和入队。')
h('2.2 服务端分工',2)
p('服务端保存账号验证材料、公开预密钥、群成员和待投递密文。路由层从已认证连接确定发件人，先校验目标与群权限，再入持久化队列。服务端持有 TLS 私钥，但不持有客户端身份私钥、会话私钥或文件密钥。')
h('2.3 资源约束',2)
p('单帧最多 8 MiB，网络出站队列最多 256 帧且累计不超过 16 MiB；GUI 命令队列上限为 64，事件队列上限为 1024，文件进度事件可合并。服务端最多保持 256 个活动连接，并定期回收已结束的连接线程。')

# 4 Protocol
h('3 协议与身份信任',page=True)
h('3.1 TLS 与账号认证',2)
p('正式连接限制为 TLS 1.3，同时验证证书链和预期主机名。演示使用含 localhost 身份的自签证书，客户端只显式信任该文件，不向系统证书库写入根证书。TLS 验证成功后进行 v2 协商，之后才允许提交注册或登录消息。')
p('注册口令经 TLS 送到服务端处理。登录使用有时限的一次性挑战，认证证明绑定随机 nonce，并同时校验客户端身份签名。错误口令、跨连接使用挑战、超时和重放均拒绝；服务端不把单纯“发出了挑战”当作认证成功。')
h('3.2 会话建立与密文绑定',2)
p('会话建立前核对用户名绑定的签名公钥与 DH 身份公钥。未知身份和身份变化均要求用户比较指纹并显式确认；随后校验签名预密钥，再调用已有 X3DH，覆盖有一次性预密钥和无一次性预密钥两条路径。')
p('v2 加密内层绑定消息 ID、发件人、收件人、群 ID 与消息种类。接收方在候选棘轮状态上解密并核对内外层路由；认证或绑定失败时不提交候选状态，也不消费乱序缓存密钥。')
table('表3.1 密码原语与用途',['原语','用途','验证依据'],[
    ['X25519 / Ed25519','密钥协商、身份及预密钥签名','RFC 7748、RFC 8032'],
    ['HKDF-SHA256','会话与棘轮密钥派生','RFC 5869'],
    ['AES-256-GCM','正文、元数据与文件块认证加密','GCM 已知测试向量'],
    ['PBKDF2-SHA256','口令派生','RFC 7914 测试向量'],
    ['SHA-256','文件完整性、指纹与交付文件摘要','边界及真实传输测试'],
], [3.4,6.4,4.863], short=(0,))
h('3.3 协议边界',2)
p('算法通过 OpenSSL 3 EVP 调用，线格式与状态持久化由本项目定义。旧 v1 入口保留供回归，正式桌面客户端采用 v2 安全路径。身份核对不能由一条“连接成功”提示代替；指纹必须通过可信渠道比较。')

# 5 Durable storage
h('4 本地存储与可靠恢复',page=True)
h('4.1 原子快照',2)
p('客户端使用 SQLite 保存账号、身份、信任记录、棘轮状态、聊天历史、outbox、文件检查点以及群快照。包含秘密的快照通过 Windows DPAPI 保护，绑定当前 Windows 用户和账号作用域；同名账号连接不同服务端时进入不同 profile。提交使用版本检查与数据库事务，冲突、损坏和保护失败均返回明确错误。')
h('4.2 发送与接收顺序',2)
p('发送顺序为：构造候选密文 → 一次提交棘轮、历史和 outbox → 网络发送。提交失败时不发送；掉线重试时复用原消息 ID 与原密文，避免再次推进会话。')
p('接收顺序为：候选解密 → 校验身份和路由 → 去重并持久化 → 发出签收并发布界面事件。若落盘失败，接收方不 ACK。服务端先持久化密文再确认入队，收到授权接收方的签收后才删除待投递项。')
table('表4.1 私聊投递状态与含义',['状态','含义','保留策略'],[
    ['待发送','本地事务已提交，等待发送或重试','保留原密文 outbox'],
    ['服务端已接收','服务端可靠队列已接纳','继续等待对端签收'],
    ['对端已签收','对端已认证解密并成功提交','移除待发项，保留历史'],
], [3.3,7.2,4.163], short=(0,))
h('4.3 重启后的行为',2)
p('启动先加载原身份、信任和会话状态，再登录与恢复未完成工作。聊天 outbox 自动补发；文件恢复重新验证源文件和已确认的接收前缀。已有账号的身份数据缺失或损坏时停止相关操作，不悄悄生成新身份替代。')
p('M5-03 在双方真实进程终止及签收丢失后重启客户端，验证原密文补投、历史只展示一条、接收棘轮不重复推进。M3-03 与 M3-04 分别验证事务故障、作用域绑定和 DPAPI 保护失败等路径。')
h('4.4 静态数据保护范围',2)
p('DPAPI 保护的是客户端快照，不防护已控制当前 Windows 用户的进程。源文件、接收完成文件以及保留的临时文件仍是普通文件；跨 Windows 用户复制数据库不等同于可恢复备份。')

# 6 Files
h('5 文件分块与续传',page=True)
p('每次传输产生独立随机 32 字节文件密钥和传输 ID，文件名、长度及 SHA-256 等元数据经会话棘轮加密发送。文件密钥不从会话根密钥提取。每个数据块为 64 KiB，密文与 16 字节认证标签分别处理，零字节及短密文均走合法边界。')
figure(ROOT/'artifacts/gui/bob-screens/file-progress.png','图5-1 接收端显示已认证的文件进度',width=13.0)
h('5.1 持久化检查点',2)
p('接收端先写 .part 并 flush，再提交检查点，最后发加密 ACK；只有完整长度与摘要验证通过才改为最终文件名。检查点保留已确认位置、前块标签、文件密钥、源摘要、接收前缀摘要和在途原密文。')
h('5.2 恢复与取消',2)
p('真实续传测试在双方已提交 128 KiB 后终止进程，重启完成 2 MiB+1 字节文件并核对摘要。源文件变化或临时前缀被改动时拒绝恢复；取消保留临时文件与记录，再次发送必须创建新传输。单文件上限 1 GiB，最多 16 个活动传输，乱序缓存最多 64 块。')

# 7 Groups
h('6 群聊与在线状态',page=True)
p('群聊采用逐成员独立会话加密。一个逻辑群消息为每位其他成员生成独立的传输 ID 和密文；服务端保存群成员元数据并校验发件人与收件人均在群中，不保存共享群密钥。群大小为 2–64 人，成员不重复，快照带递增版本。')
figure(ROOT/'artifacts/gui/alice-screens/group.png','图6-1 群消息与逐成员发送结果',width=13.0)
h('6.1 局部失败与重试',2)
p('群草稿先持久化，再逐成员保存消息 ID 或错误。失败不会抹掉其他成员的成功结果；重试只面向失败成员。如果聊天已提交而群报告提交失败，恢复通过相同逻辑 ID 找回已有消息，避免重复加密和重复展示。')
h('6.2 在线与输入提示',2)
p('登录后主动获取权威用户名单，断线清空在线和输入状态，重连后重新同步。输入提示使用单调时钟，每秒最多发送一次，3 秒后自动到期。三客户端群测试另行验证离线成员补投、独立密文和非成员拒绝。')

# 8 GUI
h('7 图形界面与线程响应',page=True)
p('界面包含账号与连接、聊天、文件和群聊四页。身份确认弹窗显示指纹，私聊展示签收状态，文件表显示经过认证的字节进度，群页保留每位收件人的排队结果。图7-1为真实双窗口验收中保存的私聊页面。')
figure(ROOT/'artifacts/gui/alice-screens/chat.png','图7-1 私聊页面与对端签收状态',width=13.0)
h('7.1 命令与事件队列',2)
p('PBKDF2、文件哈希、会话加密和数据库事务均在业务线程运行。GUI 每 40 ms 读取事件队列，更新控件不持有业务状态。关闭窗口时停止接收新命令，停止业务循环并 join 网络线程，再销毁服务与窗口对象。连接阶段的关闭可能等待当前连接超时完成。')
h('7.2 真实窗口验收',2)
alice_gui=json.loads((ROOT/'artifacts/gui/alice.json.done').read_text(encoding='utf-8-sig'))
bob_gui=json.loads((ROOT/'artifacts/gui/bob.json.done').read_text(encoding='utf-8-sig'))
p(f'M4-04 通过 Qt Test 驱动两个独立原生窗口，完成注册、双方指纹核对、互聊、1 MiB+15 字节文件进度、群发和正常关闭。归档样本中 Alice 与 Bob 的 GUI 心跳分别记录 {alice_gui["gui_heartbeats"]}、{bob_gui["gui_heartbeats"]} 次，最大间隔为 {alice_gui["max_gui_gap_ms"]}、{bob_gui["max_gui_gap_ms"]} ms；这些结果反映本机样本，不代表所有网络条件下的响应保证。')

# 9 Testing
h('8 验证结果与证据',page=True)
p('验证环境为 Windows x64，GNU 16.2.0、OpenSSL 3.6.4、SQLite 3.53.4 与 Qt 6.11.2。原项目单项 CTest 内含 50 个 doctest case。纯核心独立验收为 20 项 CTest；桌面与交付验收增加 M4-04 和 M6-01，共 22 项。主要覆盖关系见表8.1。')
table('表8.1 已完成任务与可复查证据',['任务','验证内容','证据入口'],[
    ['M0-01～03','契约、TLS 并发回归、客户端无 server 链接','framework-contracts，M0-02/03'],
    ['M1-01～03','已知向量、棘轮失败原子性、身份信任','M1-01/02/03'],
    ['M2-01～03','CA/主机名/有效期、挑战生命周期、可靠队列','M2-01/02/03'],
    ['M3-01～05','传输、账号、原子快照、DPAPI、编排线程','M3-01/02/03/04/05'],
    ['M4-01','双 TLS 文本、篡改、去重和提交故障','M4-01'],
    ['M4-02','0/1/15/16/65535/65536/65537 字节、错 ACK','M4-02'],
    ['M4-03','真实进程终止后续传、原密文与源变化拒绝','M4-03'],
    ['M4-04','双原生 GUI 身份核对、互聊、进度和关闭','M4-04；artifacts/gui'],
    ['M5-01～03','三客户端群、在线到期、聊天跨进程恢复','M5-01/02/03'],
    ['M6-01','发布包独立启动、TLS 收发、文件摘要与文档 QA','M6-01；package-smoke.json'],
], [2.8,7.5,4.363], short=(0,))
h('8.1 初始失败与修复记录',2)
p('初始基线为 43 cases 中 42 通过，注册等待失败。根因为同一 SSL 对象被不同线程并发读写，修为单网络线程拥有；后续向量和原子性用例使原测试扩展至 50 cases。文件续传测试还修复了进程提前退出丢弃最终收据的同步问题，两次原失败日志均保留。')
h('8.2 结论的适用范围',2)
p('证据覆盖本机功能、失败路径与恢复一致性。发布包测试限制子进程 PATH 仅包含程序包和 Windows，验证真实 Release 服务端认证、密文往返，以及实际 GUI 启动、响应和关闭；双 GUI 业务操作另由 M4-04 验证。未进行独立安全审计、长期公网压力测试或 Signal 互操作测试。')

# 10 Operating instructions
h('9 运行与复现',page=True)
h('9.1 免安装演示',2)
p('解压 CipherChat-Windows-x64.zip，在包目录打开 PowerShell 执行以下命令。脚本启动一台服务端和 Alice、Bob 两个窗口，证书与数据在首次运行时生成；第三个窗口可用于群聊演示。')
for line in ['.\\start-demo.ps1', '.\\start-demo.ps1 -ThirdClient', '.\\start-demo.ps1 -Port 55221 -DataDirectory D:\\CipherChatDemo2']:
    p(line,'CC Code')
p('首次在各自窗口输入口令并选择“注册并登录”，再次使用原口令选择“登录”。确认身份前，双方在账号页复制自己的指纹，通过可信渠道比较。私聊可查看签收；文件页选择联系人与源文件，恢复时选中原传输；群页输入群名及已注册成员，再创建和发送。')
p('默认数据根目录为 %LOCALAPPDATA%\\CipherChat\\demo，两个窗口使用独立目录。全部窗口关闭后脚本停止本次服务端。服务端默认监听所有 IPv4 接口，演示客户端使用 127.0.0.1；端口被占用时脚本报错，可显式换端口。')
h('9.2 手动启动',2)
p('若本机策略不允许执行脚本，可在包目录手动运行服务端，再打开两次客户端，不需要修改系统策略。服务端参数依次为端口、数据库、证书与私钥路径。')
p('.\\cipherchat-server.exe 55220 server.db cert.pem key.pem','CC Code')
p('客户端填 127.0.0.1、55220，选择服务端 cert.pem，证书主机名填 localhost，使用不同用户名与数据目录。不要把服务端 key.pem 当作客户端信任证书。')
h('9.3 源码构建',2)
p('依赖位置由 CMakePresets.json 配置。核心构建关闭 Qt；桌面发布构建使用 Release。以下命令在项目根目录运行，完整验收和依赖 PATH 见 README。')
for line in ['cmake --preset desktop-release', 'cmake --build --preset desktop-release --target', '  cipherchat-server ccfw-desktop -j 3', '.\\scripts\\package-windows.ps1', 'ctest --test-dir framework/build/framework-desktop', '  --output-on-failure --timeout 90']:
    p(line,'CC Code')
p('上方换行用于排版，带参数的构建与测试命令应各作为一条命令执行。打包脚本拒绝覆盖已有目录，重新打包可传入新的 -Destination。包内 manifest.json 保存每个文件的长度与 SHA-256，交付包不包含账号数据库或私钥。')

# 11 Boundaries and references
h('10 适用边界与改进方向',page=True)
h('10.1 当前实现边界',2)
p('系统已完成本课程设计要求的账号、私聊、文件续传、群聊、状态和恢复流程。服务端仍可见用户名、路由、群成员、在线状态、时间与大小等元数据；端到端加密不隐藏这些信息。注册口令在经验证的 TLS 连接内提交服务端，静态快照保护依赖当前 Windows 用户的 DPAPI。')
p('当前不提供多设备同步、密码找回、身份自动迁移、群成员编辑或完整公网运维能力。收到身份变化提示时需要重新核对；账号本地身份丢失时不能通过静默换钥继续使用原账号。文件取消保留临时数据，不承担自动清除用户文件的职责。')
h('10.2 后续工程方向',2)
p('后续可在保持协议版本边界的前提下增加备份与迁移设计、可管理的群成员更新、网络限速和长时间压力测试。面向公网部署前，应进一步评估身份恢复策略、服务端容量、异常连接处理、证书生命周期以及独立安全审计。上述内容属于后续工作，不计入本次已完成功能。')
h('参考文献')
references = [
    '[1] Langley A, Hamburg M, Turner S. Elliptic Curves for Security. RFC 7748, 2016. https://www.rfc-editor.org/rfc/rfc7748',
    '[2] Josefsson S, Liusvaara I. Edwards-Curve Digital Signature Algorithm. RFC 8032, 2017. https://www.rfc-editor.org/rfc/rfc8032',
    '[3] Krawczyk H, Eronen P. HMAC-based Extract-and-Expand Key Derivation Function. RFC 5869, 2010. https://www.rfc-editor.org/rfc/rfc5869',
    '[4] Percival C, Josefsson S. The scrypt Password-Based Key Derivation Function. RFC 7914, 2016, Section 11. https://www.rfc-editor.org/rfc/rfc7914',
    '[5] Marlinspike M, Perrin T. The X3DH Key Agreement Protocol; The Double Ratchet Algorithm. Signal specifications. https://signal.org/docs/',
    '[6] OpenSSL Project. OpenSSL 3 EVP Documentation. https://docs.openssl.org/3.0/man7/evp/',
    '[7] Microsoft. Data Protection API. https://learn.microsoft.com/windows/win32/secauthn/data-protection',
    '[8] CipherChat 项目. 源码、framework/tasks.json、contracts.md 与本次验收日志. 2026-09-14.',
]
for ref in references: p(ref,'CC Reference')

# Normalize direct text formatting throughout the body, tables and header/footer.
for element in doc._element.iter(qn('w:rPr')):
    color=element.find(qn('w:color'))
    if color is None: color=OxmlElement('w:color'); element.append(color)
    color.attrib.clear(); color.set(qn('w:val'),'000000')
for section in doc.sections:
    for part in (section.header,section.footer,section.first_page_header,section.first_page_footer):
        for par in part.paragraphs:
            for run in par.runs:
                run.font.name='Times New Roman'; run.font.size=Pt(10.5); run.font.color.rgb=BLACK
doc.core_properties.title='CipherChat 端到端加密聊天系统项目方案书'
doc.core_properties.subject='设计方案、实现、验证与运行说明'
doc.core_properties.author=''; doc.core_properties.last_modified_by=''
doc.core_properties.comments=''; doc.core_properties.keywords='CipherChat;课程设计;端到端加密'
doc.save(OUT)
with zipfile.ZipFile(OUT) as archive:
    archive.testzip()
record={'template_sha256':hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
        'document_sha256':hashlib.sha256(OUT.read_bytes()).hexdigest(),
        'embedded_images':len(doc.inline_shapes), 'tables':len(doc.tables),
        'paragraphs':len(doc.paragraphs), 'output':str(OUT)}
(WORK/'build-record.json').write_text(json.dumps(record,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(record,ensure_ascii=False))
